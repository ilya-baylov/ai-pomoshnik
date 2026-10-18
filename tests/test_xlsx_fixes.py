from pathlib import Path
import zipfile
import pytest
from openpyxl import Workbook
from app.ingest.loaders.tabular import XlsxLoader
from app.ingest.loaders.base import BlockKind


def save(tmp_path, rows, merges=()):
    wb=Workbook()
    ws=wb.active
    for row in rows:ws.append(row)
    for value in merges:ws.merge_cells(value)
    path=tmp_path/'test.xlsx'
    wb.save(path)
    return path


def rows(path):
    return [b.text for b in XlsxLoader().load(path).blocks if b.kind==BlockKind.TABLE_ROW]


def test_blank_days_and_exclusions_do_not_inherit_previous_master(tmp_path):
    path=save(tmp_path,[['Мастер','ПН','ВТ','Не делает'],['Анна','9-18','9-18','Наращивание'],['Ольга',None,'10-19',None]])
    result=rows(path)
    olga=next(r for r in result if 'Ольга' in r)
    assert 'ПН:' not in olga and 'Наращивание' not in olga


def test_only_explicit_vertical_merge_is_carried(tmp_path):
    path=save(tmp_path,[['Мастер','ПН','ВТ'],['Анна','9-18',None],[None,None,'10-19']],['A2:A3'])
    result=rows(path)
    assert 'Мастер: Анна; ВТ: 10-19' in result
    assert result[-1].count('9-18')==0


@pytest.mark.parametrize('data', [[],[['Колонка','Цена']], [['Колонка','Цена'],['Услуга',1500]]])
def test_short_and_empty_workbooks_do_not_crash(tmp_path,data):
    XlsxLoader().load(save(tmp_path,data))


def test_sparse_column_positions_preserved(tmp_path):
    path=save(tmp_path,[['Мастер','Салон','Процедура','Цена'],['Анна',None,'Маникюр',1500]])
    assert 'Процедура: Маникюр' in rows(path)[0]
    assert 'Салон:' not in rows(path)[0]


def test_percent_format_is_preserved(tmp_path):
    wb=Workbook();ws=wb.active
    ws.append(['Акция','Скидка']);ws.append(['Студентам',.2]);ws['B2'].number_format='0%'
    path=tmp_path/'percent.xlsx';wb.save(path)
    assert 'Скидка: 20%' in rows(path)[0]


@pytest.mark.parametrize('day,month,fmt,expected', [
    (2, 2, 'd/m', '2/2'), (3, 1, 'd/m', '3/1'),
    (3, 2, 'd/m', '3/2'), (3, 3, 'd/m', '3/3'),
    (5, 2, 'd/m', '5/2'), (5, 2, 'dd/mm', '05/02'),
    (5, 2, 'm/d', '2/5'), (5, 2, 'yyyy-mm-dd', '2024-02-05 00:00:00'),
])
def test_excel_date_display_is_preserved(tmp_path, day, month, fmt, expected):
    from datetime import datetime
    from app.ingest.loaders.xlsx_xml import read_xlsx_rows
    wb = Workbook(); ws = wb.active
    ws.append(['Мастер', 'ГРАФИК']); ws.append(['Анна', datetime(2024, month, day)])
    ws['B2'].number_format = fmt
    path = tmp_path / 'shift.xlsx'; wb.save(path)
    sheets, warnings = read_xlsx_rows(path)
    assert sheets[0]['rows'][1][1][1] == expected
    assert not warnings


def test_iso_typed_excel_shift_preserves_display(tmp_path):
    from datetime import datetime
    from app.ingest.loaders.xlsx_xml import read_xlsx_rows
    wb = Workbook(iso_dates=True); ws = wb.active
    ws.append(['Мастер', 'ГРАФИК']); ws.append(['Анна', datetime(2024, 2, 5)])
    ws['B2'].number_format = 'd/m'
    path = tmp_path / 'shift-iso.xlsx'; wb.save(path)
    sheets, warnings = read_xlsx_rows(path)
    assert sheets[0]['rows'][1][1][1] == '5/2'
    assert not warnings


def test_source_city_overrides_template_headers_and_excludes_foreign_rows(tmp_path):
    import yaml,hashlib
    path=save(tmp_path,[['№','Фамилия и имя сотрудника в Новосибирске','Филиалы в Новосибирске'],[1,'Чужой','Версаль'],[2,'Инесса','Баумана']])
    meta={'source_city':'Казань','table_title':'Услуги мастеров','source_sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'skip_rows':{'Sheet':[2]}}
    path.with_suffix('.xlsx.meta.yaml').write_text(yaml.safe_dump(meta,allow_unicode=True),encoding='utf8')
    result=rows(path)
    assert len(result)==1 and 'Город: Казань' in result[0] and 'Инесса' in result[0]
    assert 'Новосибирск' not in result[0]


def test_numbered_template_rows_are_ignored(tmp_path):
    path=save(tmp_path,[['№','Акция','Описание'],[1,'Скидка','Студентам'],[2,None,None],[3,None,None]])
    assert len(rows(path))==1


def test_explicit_header_overrides_resolve_ambiguous_city_labels(tmp_path):
    path=save(tmp_path,[['В Красноярске','Специальность','В Красноярске'],['Анна','Маникюр','Копылова']])
    path.with_suffix('.xlsx.meta.yaml').write_text('source_city: Красноярск\nheader_overrides:\n  1: Мастер\n  3: Салон / филиал\n',encoding='utf8')
    result=rows(path)[0]
    assert 'Мастер: Анна' in result and 'Салон / филиал: Копылова' in result


def test_hidden_sheets_not_used_as_current_information(tmp_path):
    wb=Workbook();ws=wb.active
    ws.append(['Название','Цена']);ws.append(['Услуга',100]);old=wb.create_sheet('Архив');old.sheet_state='hidden'
    old.append(['Название','Цена']);old.append(['Устарело',10]);path=tmp_path/'hidden.xlsx';wb.save(path)
    loaded=XlsxLoader().load(path)
    assert all('Устарело' not in b.text for b in loaded.blocks)
    assert loaded.warnings


def test_luna_uses_no_reasoning_by_default():
    from app.generation.llm import LLMChain,_request_kwargs
    params=_request_kwargs(LLMChain('https://api.openai.com/v1','test','gpt-6-luna'),'q',.2,1200)
    assert params['reasoning_effort']=='none' and 'temperature' not in params


@pytest.mark.asyncio
@pytest.mark.parametrize('question,expected', [
    ('Какие акции в Ростове?', 'Ростов'),
    ('Сравни акции в Ростове и Казани', None),
    ('Какие акции в Казани и Казань?', 'Казань'),
])
async def test_multiple_cities_do_not_silently_select_first(monkeypatch,question,expected):
    import importlib
    module=importlib.import_module('app.retrieval.search')
    monkeypatch.setattr(module,'_sections_cache',{'ростов':'Ростов','казань':'Казань'})
    assert await module.section_from_query(question)==expected


@pytest.mark.asyncio
async def test_sidecar_metadata_change_reingests_without_paid_vectors(tmp_path):
    import yaml
    from app.core.config import get_settings
    from app.storage import db
    from app.ingest.pipeline import Ingestor
    from tests.test_pipeline_e2e import FakeEmbeddings
    from sqlalchemy import text
    settings=get_settings()
    settings.embedding_model='fake-test-model'
    settings.embedding_dim=64
    await db.init_db()
    settings.knowledge_path.mkdir(parents=True,exist_ok=True)
    path=save(settings.knowledge_path,[['Мастер','Цена'],['Анна',100]])
    sidecar=path.with_suffix('.xlsx.meta.yaml')
    sidecar.write_text('tags: []\n',encoding='utf8')
    ingestor=Ingestor(provider=FakeEmbeddings())
    assert (await ingestor.run()).ok
    sidecar.write_text('tags: [current]\n',encoding='utf8')
    result=await ingestor.run()
    assert result.files_changed==1 and result.embeddings_computed==0
    async with db.session_scope() as session:
        assert (await session.execute(text('SELECT tags FROM documents'))).scalar()=='["current"]'
